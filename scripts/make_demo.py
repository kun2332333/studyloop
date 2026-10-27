"""生成一份中文样例讲义（pptx）用于试用与自测。

运行： python scripts/make_demo.py
输出： 样例讲义/记忆与学习_认知科学讲义.pptx
"""
import os
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, "vendor"))
from pptx import Presentation          # noqa: E402
from pptx.util import Inches, Pt       # noqa: E402
from pptx.dml.color import RGBColor    # noqa: E402

OUT_DIR = os.path.join(BASE, "样例讲义")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, "记忆与学习_认知科学讲义.pptx")

SLIDES = [
    ("记忆与学习：从突触到策略",
     ["一份给正在备考的你的认知科学速览",
      "本讲回答三个问题：记忆如何形成、为什么会忘、怎样学才不白学"],
     "开场先用一句话立起主线：记忆不是录音机，而是每次提取都在重建。"),
    ("记忆的三级加工模型",
     ["感觉记忆：视觉图像约 0.5 秒，声像约 2 秒，几乎全部遗忘",
      "短时记忆：容量约 7±2 个组块，保持约 20 秒，不复习立即消失",
      "长时记忆：容量近乎无限，但需要编码与巩固才能真正存进去",
      "组块化：把零散信息打包成一个组块，等于扩大了短时记忆容量"],
     "举例：手机号 138 0013 8000 分成三组，比 11 个孤立数字好记得多。"),
    ("工作记忆：学习的瓶颈",
     ["工作记忆是当前正在操作信息的临时工作台",
      "中央执行系统负责注意分配，语音环路负责语言材料，视空模板负责图像与空间",
      "工作记忆容量有限，这是认知负荷理论的出发点",
      "内在负荷来自材料难度，外在负荷来自糟糕的呈现方式，相关负荷才真正用于学习",
      "降低外在负荷：一页只讲一件事，去掉无关装饰，把步骤写成清单"],
     "所以本讲义每页只放 4~6 条要点——这不是排版偷懒，是认知负荷管理。"),
    ("记忆的生理基础",
     ["记忆的物理载体是突触连接的强度变化，也就是突触可塑性",
      "海马体负责把新经验快速写入，是陈述性记忆的中转站",
      "长时程增强（LTP）指高频刺激后突触传递效率持久提高，被看作学习的基本机制",
      "长时程抑制（LTD）则是相反方向的调节，两者共同完成精细的权重调整",
      "睡眠期间海马体会把白天的记忆重放给大脑皮层，完成系统巩固"],
     "强调：记忆是生物过程，需要蛋白质合成与睡眠，不是靠意志力硬扛。"),
    ("为什么会遗忘：遗忘曲线",
     ["艾宾浩斯遗忘曲线：新学材料在最初 24 小时内遗忘最快，之后逐渐变缓",
      "遗忘不是记忆消失，多数是提取线索丢失，也就是想不起来而不是没存进去",
      "对抗遗忘最有效的手段是间隔重复：在即将遗忘的时刻复习",
      "每次成功回忆都会把遗忘曲线整体抬高，衰减速度也随之变慢",
      "理想的复习间隔随掌握程度拉长：1 天、3 天、7 天、16 天、35 天"],
     "本软件用 FSRS 算法自动计算每张卡的复习间隔，就是你个人的遗忘曲线。"),
    ("主动回忆与测试效应",
     ["提取练习：主动回忆比重读有效得多，这叫测试效应",
      "测试本身不是评估手段，而是最强的学习手段之一",
      "重读带来熟悉感，熟悉感会被误判为掌握，这就是流畅性错觉",
      "合上书用自己的话复述，比划线高亮有效 2 倍以上",
      "难度适中的提取最有效：太容易没有增益，完全想不起来会挫败"],
     "所以复习界面默认先让你自己回忆，再翻面看答案，绝不一上来就给答案。"),
    ("睡眠、运动与记忆巩固",
     ["慢波睡眠阶段的记忆重放是巩固的关键窗口",
      "睡前 30 分钟复习当天最难的内容，利用睡眠完成巩固",
      "睡眠不足会显著降低海马体的编码能力，熬夜背书是最差的策略",
      "有氧运动提升脑源性神经营养因子（BDNF），间接支持突触可塑性",
      "分散学习优于集中突击：同样 6 小时，分成 6 天效果远好于一天"],
     "如果你连续熬夜，记得再多也只是把信息堆在短时记忆里，第二天就没了。"),
    ("高效学习策略清单",
     ["主动回忆：先自测，再查漏，最后才重读",
      "间隔重复：把复习分散到多天，而不是考前突击",
      "交错练习：不同题型、不同科目交替练，提升辨别与迁移能力",
      "精细加工：多问为什么，把新知识接到旧知识的钩子上",
      "费曼技巧：用大白话讲给外行听，讲不通的地方就是没懂的地方",
      "双编码：文字配图，图像与语言两条通路同时编码，记忆更牢"],
     "这六条正是本软件功能设计的依据：功能不是拍脑袋加的，每条都对应一个效应。"),
]


def make_chart(path):
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None
    W, H = 900, 520
    img = Image.new("RGB", (W, H), "#11141c")
    d = ImageDraw.Draw(img)
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    try:
        f = ImageFont.truetype(font_path, 22)
        fs = ImageFont.truetype(font_path, 16)
    except OSError:
        f = fs = ImageFont.load_default()

    left, right, top, bottom = 90, 60, 50, 70
    pw, ph = W - left - right, H - top - bottom
    d.line([(left, top), (left, top + ph)], fill="#3a4152", width=2)
    d.line([(left, top + ph), (left + pw, top + ph)], fill="#3a4152", width=2)
    d.text((left, 14), "复习如何抬高你的记忆保持率", font=f, fill="#e8ecf5")
    d.text((left - 60, top + 4), "100%", font=fs, fill="#8b93a7")
    d.text((left - 45, top + ph - 10), "0%", font=fs, fill="#8b93a7")
    d.text((left + pw - 44, top + ph + 14), "天", font=fs, fill="#8b93a7")

    import math

    def curve(x, k):
        return top + ph * (1 - math.exp(-k * x / 10)) * 0.98

    pts = [(left + pw * x / 100, curve(x, 4.2)) for x in range(0, 101, 2)]
    d.line(pts, fill="#5a6275", width=3)
    d.text((left + 20, top + ph - 34), "不复习：24 小时内断崖式下降", font=fs, fill="#7f879b")

    pts2 = [(left + pw * x / 100, curve(x, 0.55)) for x in range(0, 101, 2)]
    d.line(pts2, fill="#6C8CFF", width=4)
    for i, x in enumerate((8, 22, 45, 78)):
        cx, cy = left + pw * x / 100, curve(x, 0.55)
        d.ellipse([cx - 7, cy - 7, cx + 7, cy + 7], fill="#5FD3A6")
        d.text((cx - 8, cy - 34), f"{i + 1}", font=fs, fill="#5FD3A6")
    d.text((left + 20, top + 40), "间隔重复：每次提取都把曲线整体抬高", font=fs, fill="#9fb3ff")
    img.save(path)
    return path


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    chart = make_chart(os.path.join(OUT_DIR, "_chart.png"))

    for i, (title, bullets, note) in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        tb = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(1.0))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(30)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x22, 0x33)

        body = slide.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(11.4),
                                        Inches(4.6 if not (i == 4 and chart) else 2.9))
        tf = body.text_frame
        tf.word_wrap = True
        for j, b in enumerate(bullets):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            para.text = "• " + b
            para.font.size = Pt(19)
            para.space_after = Pt(10)
            para.font.color.rgb = RGBColor(0x2A, 0x30, 0x40)

        if i == 4 and chart:
            slide.shapes.add_picture(chart, Inches(1.4), Inches(4.9), width=Inches(6.2))

        slide.notes_slide.notes_text_frame.text = note

    prs.save(OUT)
    print("已生成：", OUT)


if __name__ == "__main__":
    main()