"""讲义解析：PPTX / PDF / DOCX / TXT / MD → 结构化小节（sections）。

以及"从讲义生成复习卡片"的启发式算法（无需联网、无需大模型）：
1. 断句 → 2. 全讲义 n-gram 词频找出真正的"术语" → 3. 生成三类卡：
   - 定义卡（X 是什么）  - 挖空卡（关键术语 cloze）  - 图片卡（图表页）
"""
from __future__ import annotations

import hashlib
import os
import re
import sys
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VENDOR = os.path.join(BASE_DIR, "vendor")
if os.path.isdir(VENDOR) and VENDOR not in sys.path:
    sys.path.insert(0, VENDOR)

from core import db  # noqa: E402

CJK = r"\u4e00-\u9fff"
SENT_SPLIT = re.compile(r"(?<=[。！？!?；;\n])")
BULLET = re.compile(r"^[\s•·▪◦\-\*\u25cf\u25aa]+")
NOISE = [
    re.compile(r"^[\d\s\-–—.、/]+$"),
    re.compile(r"^(谢谢|感谢|the end|thanks|本节完|下节预告)", re.I),
    re.compile(r"^\s*(©|copyright|copyrighted)", re.I),
    re.compile(r"https?://"),
]

# 高频功能字：术语里出现这些字，基本可以判定不是术语
FUNC_CHARS = set("的了是在和也就都而及对把被让给从为以这那个们你我他它很更最会能要不没可还但或与其之所着过又再才只等将已并且如果因为没有一些什么怎样如何所每各该单")

# 词里出现这些片段，多半是把动词短语当成术语了
BAD_FRAGS = ("来自", "进行", "包括", "属于", "意味", "导致", "产生", "提供", "依赖",
             "取决", "需要", "可以", "我们", "你们", "他们", "什么", "如何", "因为",
             "所以", "但是", "并且", "或者", "以及", "这是", "就是", "不是", "没有",
             "之间", "之后", "之前", "通过", "的是", "一个", "一种", "这个", "那个",
             "的方", "的时", "的是")
LABEL_STOP = {"举例", "例如", "强调", "注意", "定义", "小结", "总结", "提示", "练习",
              "问题", "答案", "课程", "目录", "引言", "结论", "本讲", "本节", "本课",
              "本文", "本章", "本软件"}

STOP_TERMS = {
    "我们", "你们", "他们", "大家", "这个", "那个", "这些", "那些", "一个", "一种",
    "可以", "如何", "什么", "因为", "所以", "但是", "如果", "然后", "以及", "并且",
    "这样", "那样", "一些", "很多", "非常", "已经", "还是", "就是", "不是", "没有",
    "进行", "通过", "使用", "需要", "能够", "可能", "应该", "问题", "内容", "方面",
    "情况", "过程", "时候", "地方", "东西", "部分", "方法", "方式", "结果", "时候",
    # 泛化词：在本讲义里会反复出现，但作为考点毫无区分度
    "记忆", "学习", "复习", "考试", "知识", "信息", "经验", "能力", "效果", "策略",
    "自己", "同时", "负责", "分成", "成为", "包括", "认为", "表示", "说明", "解释",
    "影响", "作用", "关系", "特点", "意义", "提高", "降低", "增加", "减少", "建立",
    "核心", "主要", "重要", "基本", "有效", "真正", "一定", "不同", "相同", "大量",
    "本软件", "本讲", "本节", "本课", "本文", "本章", "举例", "强调", "说明", "例子",
}

# 动词/副词结尾的候选，多半不是术语
BAD_SUFFIX = ("的", "了", "是", "在", "和", "与", "就", "也", "都", "还", "会", "能",
              "要", "把", "被", "让", "给", "从", "对", "而", "则", "并", "或", "等",
              "着", "过", "地", "得", "很", "更", "最",
              # 单字动词/趋向词结尾，几乎总是不是术语
              "来", "去", "到", "做", "用", "有", "无", "进", "出", "成", "好", "大",
              "小", "多", "少", "上", "下", "里", "内", "外", "前", "后", "我", "你")

DEF_PAT = re.compile(
    r"^(.{2,18}?)(?:指的是|是指|定义为|被称为|称为|叫做|就是|是)(.+)$")


# ---------------------------------------------------------------- 基础工具

def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = []
    for line in text.split("\n"):
        line = BULLET.sub("", line).strip()
        line = re.sub(r"^\d+[.、)]\s*", "", line)
        if line:
            lines.append(line)
    return "\n".join(lines).strip()


def is_noise(text: str) -> bool:
    t = text.strip()
    if len(t) < 4:
        return True
    return any(p.search(t) for p in NOISE)


def split_sentences(text: str) -> list[str]:
    out: list[str] = []
    for raw in SENT_SPLIT.split(text or ""):
        s = raw.strip()
        if not s or is_noise(s):
            continue
        if len(s) > 220:  # 长句再按逗号切
            out.extend(x.strip() for x in re.split(r"(?<=[，,])", s) if len(x.strip()) >= 8)
        else:
            out.append(s)
    seen, res = set(), []
    for s in out:
        key = re.sub(r"\s+", "", s)
        if key in seen:
            continue
        seen.add(key)
        res.append(s)
    return res


def _save_media(blob: bytes, ext: str) -> str:
    ext = (ext or "png").lower().lstrip(".")
    if ext in ("jpeg", "jpg"):
        ext = "jpg"
    if ext not in ("png", "jpg", "gif", "bmp", "webp"):
        ext = "png"
    h = hashlib.sha1(blob).hexdigest()[:16]
    name = f"{h}.{ext}"
    path = os.path.join(db.MEDIA_DIR, name)
    if not os.path.exists(path):
        with open(path, "wb") as f:
            f.write(blob)
    return f"media/{name}"


# ---------------------------------------------------------------- 词频术语挖掘

def _ngram_counts(texts: list[str], lo: int = 2, hi: int = 6) -> dict[str, int]:
    cnt: dict[str, int] = {}
    for text in texts:
        for chunk in re.split(r"[^\u4e00-\u9fff]+", text or ""):
            n = len(chunk)
            if n < lo:
                continue
            for size in range(lo, min(hi, n) + 1):
                for i in range(n - size + 1):
                    cnt[chunk[i:i + size]] = cnt.get(chunk[i:i + size], 0) + 1
    return cnt


def _harvest(texts: list[str]) -> dict[str, float]:
    """补充挖掘：只出现一次、但结构上明确是术语的片段。

    - "组块化：把零散信息…"  冒号前的标签
    - "长时程增强（LTP）"    带英文缩写的术语
    - "记忆巩固是指…"       定义句的主语
    """
    found: dict[str, float] = {}

    def ok(t: str) -> bool:
        return (2 <= len(t) <= 10 and t not in LABEL_STOP
                and not any(ch in FUNC_CHARS for ch in t)
                and not any(f in t for f in BAD_FRAGS)
                and t[-1] not in BAD_SUFFIX)

    for text in texts:
        for line in (text or "").split("\n"):
            m = re.match(rf"^([{CJK}]{{2,10}})[：:]", line)
            if m and ok(m.group(1)):
                found[m.group(1)] = max(found.get(m.group(1), 0), 3.0)
        for m in re.finditer(rf"([{CJK}]{{3,10}})[（(]([A-Za-z]{{2,8}})[)）]", text or ""):
            if ok(m.group(1)):
                found[m.group(1)] = max(found.get(m.group(1), 0), 4.5)
        for m in re.finditer(rf"([{CJK}]{{3,10}})(?:是指|指的是|定义为|称为|叫做)", text or ""):
            if ok(m.group(1)):
                found[m.group(1)] = max(found.get(m.group(1), 0), 4.0)
    return found


def key_terms(texts: list[str], top: int = 80) -> dict[str, float]:
    """术语发现 = n-gram 词频 + 结构挖掘。

    真术语会在讲义里反复出现（工作记忆、突触可塑性…），
    随机字串只出现一次，因此 freq>=2 是很有力的过滤器；
    只出现一次但结构明确的（带缩写、冒号标签、定义主语）由 _harvest 补上。
    """
    cnt = _ngram_counts(texts)
    scored: dict[str, float] = {}
    for term, c in cnt.items():
        if c < 2:
            continue
        if term in STOP_TERMS or any(ch in FUNC_CHARS for ch in term):
            continue
        if any(f in term for f in BAD_FRAGS):
            continue
        if term[-1] in BAD_SUFFIX or term[0] in "的了和与而则并就都也还被把给从对":
            continue
        if re.search(r"(..)\1{2,}", term):
            continue
        n = len(term)
        if n == 2:
            if c < 4:
                continue
            length_bonus = 0.45
        elif n == 3:
            length_bonus = 1.0
        elif n == 4:
            length_bonus = 1.35
        elif n == 5:
            length_bonus = 1.2
        else:  # 6 字以上多半是摘抄的半句话
            length_bonus = 0.7
        scored[term] = c * length_bonus
    # 去掉被更长、同样高频的术语包含的候选（"工作记忆" 压过 "工作"）
    items = sorted(scored.items(), key=lambda kv: (-kv[1], -len(kv[0])))
    keep: dict[str, float] = {}
    for term, score in items:
        if any(term != other and term in other and cnt[other] >= cnt[term] * 0.5
               for other in keep):
            continue
        keep[term] = score

    for term, score in _harvest(texts).items():
        keep[term] = max(keep.get(term, 0), score * 1.6)
    return dict(sorted(keep.items(), key=lambda kv: -kv[1])[:top])


def _pick_term(sentence: str, terms: dict[str, float],
               usage: dict[str, int] | None = None, cap: int = 2) -> str | None:
    """在句子里挑一个"最值得挖空"的术语：优先长、优先高频、还没被用烂。"""
    best, best_key = None, (-1, -1.0)
    for term, score in terms.items():
        if term not in sentence:
            continue
        if usage and usage.get(term, 0) >= cap:
            continue
        if len(term) >= len(sentence) - 1:      # 整句就是一个词，挖了没上下文
            continue
        key = (min(len(term), 5), round(score, 2))
        if key > best_key:
            best, best_key = term, key
    return best


def _quoted(sentence: str) -> list[str]:
    return [m for m in re.findall(r"[「“\"『]([^」”\"』]{2,16})[」”\"』]", sentence)]


# ---------------------------------------------------------------- 卡片生成

def generate_cards(sections: list[dict], max_per_section: int = 6) -> list[dict]:
    texts = [s.get("text", "") for s in sections]
    terms = key_terms(texts)
    usage: dict[str, int] = {}
    cards: list[dict] = []
    seen: set[str] = set()

    def add(deck_section, kind, front, back, note="", image=""):
        norm = re.sub(r"[\s，。、；：？！,.:;?！（）()【】\[\]「」“”\"'…—\-_]", "", front)
        if not norm or norm in seen:
            return False
        seen.add(norm)
        cards.append({
            "section_id": deck_section["id"], "idx": deck_section["idx"],
            "kind": kind, "front": front, "back": back, "note": note,
            "image_path": image or deck_section.get("image_path", ""),
        })
        return True

    for sec in sections:
        text = sec.get("text", "")
        title = sec.get("title", "") or f"第{sec.get('idx', 0) + 1}节"
        # 封面页 / 目录页文字太少，生成出来的卡没有信息量
        if len(text) < 45:
            continue
        sents = split_sentences(text)
        produced = 0
        def_terms: set[str] = set()

        # 1) 定义卡：最接近"主动回忆"的形态
        for s in sents:
            if produced >= max_per_section:
                break
            m = DEF_PAT.match(s)
            if not m:
                continue
            subject, body = m.group(1).strip(), m.group(2).strip()
            if not (2 <= len(subject) <= 14) or len(body) < 6:
                continue
            if any(ch in FUNC_CHARS or ch in "：:则就也都还便即" for ch in subject):
                continue
            if subject[-1] in BAD_SUFFIX or subject[-1].isdigit():
                continue
            if any(ch in body for ch in "你我他你"):
                continue
            if add(sec, "basic", f"「{subject}」是什么？", body, note=s[:180]):
                produced += 1
                def_terms.add(subject)

        # 2) 挖空卡：术语优先，其次引号内内容
        for s in sents:
            if produced >= max_per_section:
                break
            if len(s) < 14 or len(s) > 170:
                continue
            if s.startswith(title) and len(s) < len(title) + 4:
                continue                       # 与标题重复的句子
            term = _pick_term(s, terms, usage, cap=2)
            if term in def_terms:          # 同一节里已经问过这个术语，别再挖空
                term = None
            if not term:
                q = _quoted(s)
                term = q[0] if q and len(q[0]) >= 2 else None
            if not term:
                continue
            pos = s.find(term)
            if pos < 0 or (pos == 0 and len(s) - len(term) < 6):
                continue
            masked = s[:pos] + "＿＿＿＿" + s[pos + len(term):]
            usage[term] = usage.get(term, 0) + 1
            if add(sec, "cloze", masked, term, note=s[:200]):
                produced += 1

        # 3) 图片卡：文字很少但有大图（图表 / 示意图）
        if produced == 0 and sec.get("image_path") and len(text) < 200:
            add(sec, "image", f"看图回忆：{title} 讲了什么？",
                text[:300] or "（图示内容）", image=sec["image_path"])
            produced += 1

        # 4) 兜底：整节一个"概述卡"
        if produced == 0 and sents and len(text) >= 60:
            body = " ".join(sents[:3])[:220]
            add(sec, "basic", f"{title} 的核心内容是？", body)
    return cards


# ---------------------------------------------------------------- PPTX

def extract_pptx(path: str) -> dict:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:  # pragma: no cover
        raise ValueError("缺少 PPTX 解析库 python-pptx：请运行 "
                         "python -m pip install python-pptx 后重启软件") from e

    prs = Presentation(path)
    sections = []
    for i, slide in enumerate(prs.slides):
        title = ""
        chunks: list[str] = []
        image_rel = ""
        best_area = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if not txt:
                    continue
                if shape.is_placeholder and shape.placeholder_format.idx == 0 and not title:
                    title = txt.split("\n")[0][:60]
                else:
                    chunks.append(txt)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    area = (shape.width or 0) * (shape.height or 0)
                    if area > best_area and (shape.width or 0) > 1200000:
                        best_area = area
                        image_rel = _save_media(shape.image.blob, shape.image.ext)
                except Exception:
                    pass
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        body = clean_text("\n".join(chunks))
        if notes and notes not in body:
            body = (body + "\n" + clean_text(notes)).strip()
        if not title:
            first = (body.split("\n")[0] if body else "").strip()
            title = first[:60] if first and len(first) <= 60 else f"第 {i + 1} 页"
        elif body.startswith(title):
            body = body[len(title):].strip()      # 标题在正文里重复出现，去掉
        sections.append({
            "idx": i, "title": title, "text": body,
            "image_path": image_rel,
        })
    return {"source_type": "pptx", "page_count": len(sections), "sections": sections}


# ---------------------------------------------------------------- PDF

def extract_pdf(path: str) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ValueError("缺少 PDF 解析库 pypdf：请确认程序目录里的 vendor 文件夹完整，"
                         "或运行 python -m pip install pypdf") from e

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise ValueError("PDF 已加密，无法读取") from e

    sections = []
    for i, page in enumerate(reader.pages):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception:
            text = ""
        image_rel = ""
        try:
            imgs = sorted(getattr(page, "images", []), key=lambda im: -len(im.data))
            if imgs and len(imgs[0].data) > 24000:
                ext = os.path.splitext(getattr(imgs[0], "name", "x.png"))[1] or ".png"
                image_rel = _save_media(imgs[0].data, ext)
        except Exception:
            pass
        first = text.split("\n")[0].strip() if text else ""
        title = first[:60] if first and len(first) <= 60 else f"第 {i + 1} 页"
        sections.append({"idx": i, "title": title, "text": text, "image_path": image_rel})

    if not any(s["text"] for s in sections):
        raise ValueError("这份 PDF 没有可提取的文字（可能是扫描件/图片版），"
                         "请换用可复制文字的 PDF，或先把 PPT 导出为 PDF 再上传")
    return {"source_type": "pdf", "page_count": len(sections), "sections": sections}


# ---------------------------------------------------------------- DOCX / 纯文本

def extract_docx(path: str) -> dict:
    try:
        import docx  # type: ignore
    except ImportError as e:
        raise ValueError("缺少 docx 解析库 python-docx") from e
    doc = docx.Document(path)
    sections = []
    buf: list[str] = []
    title = ""
    idx = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (p.style.name or "").lower()
        if ("heading" in style or "标题" in (p.style.name or "")) and buf:
            sections.append({"idx": idx, "title": title or f"第 {idx + 1} 节",
                             "text": clean_text("\n".join(buf)), "image_path": ""})
            idx += 1
            buf, title = [], txt
        elif "heading" in style or "标题" in (p.style.name or ""):
            title = txt
        else:
            buf.append(txt)
    if buf:
        sections.append({"idx": idx, "title": title or f"第 {idx + 1} 节",
                         "text": clean_text("\n".join(buf)), "image_path": ""})
    return {"source_type": "docx", "page_count": len(sections), "sections": sections}


def extract_text_file(path: str) -> dict:
    raw = open(path, "rb").read()
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", "replace")
    blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
    sections = []
    for i, b in enumerate(blocks):
        first = b.strip().split("\n")[0][:60]
        sections.append({"idx": i, "title": first, "text": clean_text(b), "image_path": ""})
    return {"source_type": "txt", "page_count": len(sections), "sections": sections}


def extract(path: str, filename: str) -> dict:
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext == "pptx":
        return extract_pptx(path)
    if ext == "pdf":
        return extract_pdf(path)
    if ext in ("docx",):
        return extract_docx(path)
    if ext in ("txt", "md", "markdown"):
        return extract_text_file(path)
    if ext in ("ppt", "doc"):
        raise ValueError(f"暂不支持旧版 .{ext} 格式，请另存为 .pptx / .pdf 后上传")
    raise ValueError(f"不支持的文件类型 .{ext}")